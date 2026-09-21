"""Phase 8 - Build the Response to Reviewers document.

Generates the document in the exact structure of the official IEEE Access
template ("Reviewer#N, Concern # M" / "Author response:" / "Author action:"),
driven by reviewer_comments.json so that no concern can be silently dropped and
the verbatim quotes cannot drift from the decision letter.

Responses live in 07_response/responses.json, keyed by comment id:

    {
      "R1.1": {
        "response": "...",
        "action": "...",
        "location": "Fig. 2; Sec. V-B, p. 7"
      },
      ...
    }

Any concern without an entry is emitted with a conspicuous TO BE WRITTEN marker
and counted in the coverage report, so an incomplete document is obvious at a
glance rather than discovered by the editor.

Usage:
    python build_response.py            # writes .docx and .md
    python build_response.py --check    # coverage report only
"""
from __future__ import annotations

import json
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from paths import MATRIX, RESPONSE  # noqa: E402

TITLE = ("XAI-SDN: An explainable entropy-guided machine learning framework "
         "for real-time DDoS detection in software defined networks")
MSID = "Access-2026-39885"

PREAMBLE = f"""Original Manuscript ID: {MSID}

Original Article Title: "{TITLE}"

To: IEEE Access Editor

Re: Response to reviewers

Dear Editor,

Thank you for allowing a resubmission of our manuscript, with an opportunity to
address the reviewers' comments.

We are uploading (a) our point-by-point response to the comments (below)
(response to reviewers, under "Author's Response Files"), (b) an updated
manuscript with yellow highlighting indicating changes (as "Highlighted PDF"),
and (c) a clean updated manuscript without highlights ("Main Manuscript").

Best regards,

Adeel Ahmad et al.
"""

CLOSING = """Note: References suggested by reviewers were evaluated for relevance to this
article. Those that strengthen the work have been added and are discussed in the
text rather than merely cited; where a suggested reference did not bear on a
specific argument, it was not added, in line with the Editor's guidance.
"""


REVIEWER_THANKS = {
    1: ("We thank Reviewer 1 for a careful and constructive reading. Several of "
        "the issues raised, in particular the broken equation reference and the "
        "hardware asymmetry in the timing table, pointed to real defects that we "
        "had not detected ourselves."),
    2: ("We thank Reviewer 2 for a thorough and generous review. The comments on "
        "over-claiming and on terminology were well founded, and acting on them "
        "improved the manuscript considerably."),
    3: ("We thank Reviewer 3 for questions that went to the heart of the work. "
        "The request to remove all suspected leaking features, and the questions "
        "about control-channel capacity and flow-table occupancy, led directly to "
        "three new experiments and to findings we consider among the most useful "
        "in the revised paper."),
    4: ("We thank Reviewer 4 for a detailed and practical review. The leakage "
        "audit, the cross-dataset validation and the live testbed requested here "
        "are all now part of the paper, and the resulting evidence changed several "
        "of our conclusions."),
    5: ("We thank Reviewer 5 for a precise and helpful checklist. Every item has "
        "been addressed, and the structural additions in particular have made the "
        "paper easier to follow."),
    6: ("We thank Reviewer 6 for a demanding review. The concerns raised here were "
        "substantially correct, and pursuing them uncovered an error in our own "
        "evaluation protocol that we would otherwise have carried into print. We "
        "are genuinely grateful for the rigor of this reading."),
    7: ("We thank Reviewer 7 for a concise review that identified several real "
        "weaknesses. The questions about novelty, feature detail and evaluation "
        "protocol were fair, and answering them properly reshaped the "
        "contributions we claim."),
}


def load() -> tuple[dict, dict]:
    comments = json.loads((MATRIX / "reviewer_comments.json").read_text(encoding="utf-8"))
    responses: dict = {}
    for rf in sorted(RESPONSE.glob("responses*.json")):
        responses.update(json.loads(rf.read_text(encoding="utf-8")))
    return comments, responses


def coverage(comments: dict, responses: dict) -> dict:
    ids = [c["id"] for c in comments["comments"]]
    have = [i for i in ids if i in responses
            and responses[i].get("response") and responses[i].get("action")]
    missing = [i for i in ids if i not in have]
    no_loc = [i for i in have if not responses[i].get("location")]
    return {"total": len(ids), "answered": len(have), "missing": missing,
            "answered_without_location": no_loc,
            "complete": not missing and not no_loc}


def build_markdown(comments: dict, responses: dict) -> str:
    out = [PREAMBLE, ""]
    for r in range(1, comments["n_reviewers"] + 1):
        rows = [c for c in comments["comments"] if c["reviewer"] == r]
        if not rows:
            continue
        out.append(f"\n## Reviewer #{r}\n")
        if REVIEWER_THANKS.get(r):
            out.append(REVIEWER_THANKS[r])
            out.append("")
        for n, c in enumerate(rows, 1):
            resp = responses.get(c["id"], {})
            out.append(f"**Reviewer#{r}, Concern # {n}:** {c['verbatim']}")
            out.append("")
            out.append("**Author response:** "
                       + resp.get("response", "*** TO BE WRITTEN ***"))
            out.append("")
            action = resp.get("action", "*** TO BE WRITTEN ***")
            loc = resp.get("location")
            out.append("**Author action:** " + action
                       + (f" **Location in the revised manuscript:** {loc}" if loc else ""))
            out.append("")
    out.append(CLOSING)
    return "\n".join(out)


def build_docx(comments: dict, responses: dict, path: Path) -> None:
    import docx
    from docx.shared import Pt

    d = docx.Document()
    style = d.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(11)

    for para in PREAMBLE.strip().split("\n\n"):
        d.add_paragraph(para.strip())

    for r in range(1, comments["n_reviewers"] + 1):
        rows = [c for c in comments["comments"] if c["reviewer"] == r]
        if not rows:
            continue
        d.add_paragraph()
        h = d.add_paragraph()
        run = h.add_run(f"Reviewer #{r}")
        run.bold = True
        run.font.size = Pt(13)
        if REVIEWER_THANKS.get(r):
            d.add_paragraph(REVIEWER_THANKS[r])

        for n, c in enumerate(rows, 1):
            resp = responses.get(c["id"], {})
            p = d.add_paragraph()
            run = p.add_run(f"Reviewer#{r}, Concern # {n}: ")
            run.bold = True
            p.add_run(c["verbatim"])

            p = d.add_paragraph()
            run = p.add_run("Author response: ")
            run.bold = True
            p.add_run(resp.get("response", "*** TO BE WRITTEN ***"))

            p = d.add_paragraph()
            run = p.add_run("Author action: ")
            run.bold = True
            p.add_run(resp.get("action", "*** TO BE WRITTEN ***"))
            if resp.get("location"):
                run = p.add_run("  Location in the revised manuscript: ")
                run.bold = True
                p.add_run(resp["location"])
            d.add_paragraph()

    d.add_paragraph(CLOSING.strip())
    d.save(path)


def main(argv: list[str]) -> int:
    comments, responses = load()
    cov = coverage(comments, responses)
    RESPONSE.mkdir(parents=True, exist_ok=True)
    (RESPONSE / "coverage.json").write_text(json.dumps(cov, indent=2),
                                            encoding="utf-8")
    print(f"concerns: {cov['total']}  answered: {cov['answered']}  "
          f"complete: {cov['complete']}")
    if cov["missing"]:
        print(f"  missing responses ({len(cov['missing'])}): "
              f"{', '.join(cov['missing'])}")
    if cov["answered_without_location"]:
        print(f"  answered but no location: "
              f"{', '.join(cov['answered_without_location'])}")
    if "--check" in argv:
        return 0 if cov["complete"] else 1

    md = RESPONSE / "Response_to_Reviewers.md"
    md.write_text(build_markdown(comments, responses), encoding="utf-8")
    docx_path = RESPONSE / "Response_to_Reviewers.docx"
    build_docx(comments, responses, docx_path)
    print(f"wrote {md}\nwrote {docx_path}")
    return 0 if cov["complete"] else 1


if __name__ == "__main__":
    sys.exit(main(sys.argv[1:]))
